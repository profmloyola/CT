"""R2.0 (plan v1.0): incidencias como datos (tools/r2_incidencias.json), clases A a E, hoja .xlsx (export e import) y su efecto en auditoria --v1.
Sin navegador. Siempre sobre una copia del proyecto; lo único que se lee del original es el estado real del bloque 1 (que sus `cambios` coincidan con los datos)."""
import json, os, shutil, subprocess, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ok_all = True
def check(cond, msg, extra=''):
    global ok_all
    ok_all = ok_all and bool(cond)
    print(('PASS ' if cond else 'FAIL ') + msg, extra if not cond else '')
def jl(p): return json.load(open(p, encoding='utf-8'))

# ---- el bloque 1 real: sus cambios propuestos parten de lo que hoy dicen los datos ----
d = jl(os.path.join(ROOT, 'data', 'all.json'))
reg = {x['id']: x for k in ('works', 'designers') for x in d[k]}
real = jl(os.path.join(ROOT, 'tools', 'r2_incidencias.json'))['incidencias']
b1 = [i for i in real if i.get('bloque') == 1]
check(len(b1) == 40, 'bloque 1 real: 40 filas', len(b1))
bad = []
for i in real:
    for c in i.get('cambios', []):
        x = reg.get(c['id'])
        t = x if x and c['idioma'] in ('en', '-') else (x or {}).get('es', {})
        v = t.get(c['campo']) if x else None
        ok = x is not None and ((c['de'] in v) if isinstance(v, str) and c['de'] not in ('', None) and c['campo'] in ('more', 'dates') else (v == c['de'] or (c['de'] in ('', None) and v in ('', None))))
        if not ok: bad.append((i['id'], c['id'], c['campo'], v))
check(not bad, 'los `cambios` de las incidencias coinciden con el valor actual de los datos', bad)
check(all(i['decision'] in ('', 'Aceptar', 'Rechazar', 'Otro') for i in real) and all(i['estado'] in ('abierta', 'decidida', 'aplicada', 'cerrada') for i in real), 'estados y decisiones válidos')
check(all(i['estado'] == 'abierta' and i['decision'] in ('', 'Aceptar') for i in real), 'nada decidido ni aprobado todavía (solo los «Aceptar» por defecto de la clase A)')
a_rows = [i for i in b1 if i['clase'] == 'A']
check(all((i['decision'] == 'Aceptar') == bool(i['cambios']) for i in a_rows), 'clase A: «Aceptar» por defecto solo cuando hay una corrección concreta')

tmp = tempfile.mkdtemp(prefix='lhd-inc-')
C = os.path.join(tmp, 'lhd')
shutil.copytree(ROOT, C, ignore=shutil.ignore_patterns('__pycache__', '.git', 'e8_img', 'e7_work', 'e8_work'))
def run(*a):
    return subprocess.run([sys.executable] + list(a), cwd=C, capture_output=True, text=True)
JP = os.path.join(C, 'tools', 'r2_incidencias.json')
X = 'tools/pipeline/incidencias_xlsx.py'
W = os.path.join(C, 'tools', 'pipeline', 'r2_work')
try:
    # ---- generación: idempotente, conserva la curaduría, suma lo nuevo ----
    before = jl(JP)
    r = run('tools/pipeline/r2_incidencias.py')
    after = jl(JP)
    check(r.returncode == 0 and '0 nuevas' in r.stdout and after == before, 'r2_incidencias.py es idempotente y no pisa la curaduría', r.stdout)
    fx = [
        {'id': 'dada', 'issues': ['Quité «x»: ninguna fuente lo confirma.']},                                                       # C
        {'id': 'neoclassicism', 'issues': ['El campo status no se tocó.']},                                                          # D
        {'id': 'cubism-test', 'issues': ['Las fuentes difieren en la fecha: 1851 vs 1852; no se cambió la ficha (campo dates).']},    # B
        {'id': 'scandinavian-design', 'issues': ['DATO DE LA FICHA DESACTUALIZADO: dates dice b. 1929; Britannica y el MoMA dan 2025.']},   # A
        {'ctx': 'ctx-enlightenment', 'item': 'baskerville-virgil-1757', 'drop': True, 'issues': ['Ninguna fuente sostiene el mecanismo.']},   # E
    ]
    json.dump(fx, open(os.path.join(W, 'out_tX_1.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    r = run('tools/pipeline/r2_incidencias.py')
    new = [i for i in jl(JP)['incidencias'] if i['origen'] == 'out_tX_1.json']
    clases = {i['elemento']: i['clase'] for i in new}
    check(len(new) == 5 and '5 nuevas' in r.stdout, 'agrega las incidencias nuevas de una salida', r.stdout)
    check(clases == {'dada': 'C', 'neoclassicism': 'D', 'cubism-test': 'B', 'scandinavian-design': 'A', 'ctx-enlightenment~baskerville-virgil-1757': 'E'}, 'clasifica A, B, C, D y E por palabras clave', clases)
    check(all(i['id'].startswith('INC-') and i['estado'] == 'abierta' and i['clase_origen'] == 'auto' for i in new) and new[0]['id'] == 'INC-0718', 'ids correlativos (INC-0718…), estado abierta, clase automática', [i['id'] for i in new])
    check(run('tools/pipeline/r2_incidencias.py').stdout.count('0 nuevas') == 1, 'segunda corrida: 0 nuevas')
    os.remove(os.path.join(W, 'out_tX_1.json'))
    r = run('tools/pipeline/r2_incidencias.py')
    check(len(jl(JP)['incidencias']) == 722, 'borrar la salida original no borra las incidencias (el JSON es la fuente de verdad)')
    md = open(os.path.join(C, 'tools', 'R2-INCIDENCIAS.md'), encoding='utf-8').read()
    check('INC-0001' in md and '**C**' in md, 'R2-INCIDENCIAS.md se genera desde el JSON')

    # ---- hoja: exportar el bloque siguiente ----
    from openpyxl import load_workbook
    out = os.path.join(tmp, 'b2.xlsx')
    r = run(X, 'export', out, '--siguiente', '--n', '6')
    check(r.returncode == 0 and 'Bloque 2: 6 filas' in r.stdout, 'export --siguiente arma el bloque 2 con 6 filas', r.stdout + r.stderr)
    wb = load_workbook(out)
    ws = wb['Bloque 2']
    head = [c.value for c in ws[1]]
    check(head == ['n°', 'id', 'tipo', 'campo', 'incidencia', 'valor actual', 'propuesta', 'fuentes', 'clase', 'recomendación', 'decisión', 'valor final', 'nota'], 'columnas de la hoja', head)
    rows = [[c.value for c in r_] for r_ in ws.iter_rows(min_row=2)]
    cl = [r_[8] for r_ in rows]
    check(cl == sorted(cl, key=lambda c: 'ABDEC'.index(c)), 'filas ordenadas de la clase A a la C', cl)
    check(ws.data_validations.dataValidation and 'Aceptar' in ws.data_validations.dataValidation[0].formula1, 'la columna «decisión» tiene lista desplegable')
    check('Instrucciones' in wb.sheetnames, 'hoja de instrucciones')
    ids2 = [r_[0] for r_ in rows]
    check(all(i['bloque'] == 2 for i in jl(JP)['incidencias'] if i['id'] in ids2), 'el bloque queda asignado en el JSON')
    r = run(X, 'export', os.path.join(tmp, 'b2b.xlsx'), '--siguiente', '--n', '2')
    check('Bloque 3' in r.stdout, 'el siguiente export abre el bloque 3 (no repite filas)')
    out2 = os.path.join(tmp, 'b2again.xlsx')
    run(X, 'export', out2, '--bloque', '2')
    ws2 = load_workbook(out2)['Bloque 2']
    check([r_[0].value for r_ in ws2.iter_rows(min_row=2)] == ids2, 'export --bloque N repite el mismo bloque')
    o1 = os.path.join(tmp, 'b1.xlsx')
    run(X, 'export', o1, '--bloque', '1')
    w1 = load_workbook(o1)['Bloque 1']
    r1 = {r_[0].value: [c.value for c in r_] for r_ in w1.iter_rows(min_row=2)}
    cl1 = [r_[8].value for r_ in w1.iter_rows(min_row=2)]
    check(cl1 == sorted(cl1, key=lambda c: 'ABDEC'.index(c)) and cl1[:4] == ['A'] * 4, 'bloque 1: la clase A va primero (4 filas A, luego B y D)')
    check(len(r1) == 40 and r1['INC-0659'][10] == 'Aceptar' and not r1['INC-0135'][10], 'bloque 1: la clase A trae «Aceptar», salvo la de Fortuny (sin corrección mecánica)')

    # ---- import ----
    h0 = open(JP, 'rb').read()
    wb = load_workbook(out); ws = wb['Bloque 2']
    ws['K2'], ws['M2'] = 'Aceptar', 'ok'
    ws['K3'], ws['L3'] = 'Otro', 'valor distinto'
    ws['K4'] = 'rechazar'
    bad1 = os.path.join(tmp, 'bad1.xlsx'); wb.save(bad1)
    r = run(X, 'import', bad1, '--dry')
    check(r.returncode == 0 and '3 decididas' in r.stdout and open(JP, 'rb').read() == h0, 'import --dry cuenta y no escribe', r.stdout)
    r = run(X, 'import', bad1)
    st = {i['id']: i for i in jl(JP)['incidencias']}
    check(r.returncode == 0 and st[ids2[0]]['estado'] == 'decidida' and st[ids2[0]]['decision'] == 'Aceptar' and st[ids2[0]]['nota'] == 'ok', 'import: Aceptar → decidida y nota', r.stdout)
    check(st[ids2[1]]['decision'] == 'Otro' and st[ids2[1]]['valor_final'] == 'valor distinto', 'import: Otro con valor final')
    check(st[ids2[2]]['decision'] == 'Rechazar', 'import acepta la decisión sin importar mayúsculas')
    check(st[ids2[3]]['estado'] == 'abierta', 'sin decisión sigue abierta')
    h1 = open(JP, 'rb').read()
    r = run(X, 'import', bad1)
    check(open(JP, 'rb').read() == h1, 'import es idempotente')
    for name, edit, needle in (('Otro sin valor final', lambda w: w.__setitem__('K5', 'Otro'), 'necesita'),
                                ('decisión inválida', lambda w: w.__setitem__('K5', 'Quizás'), 'decisión inválida'),
                                ('id cambiado', lambda w: w.__setitem__('B5', 'otra-cosa'), 'no coincide'),
                                ('n° desconocido', lambda w: w.__setitem__('A5', 'INC-9999'), 'desconocido')):
        wbx = load_workbook(out); edit(wbx['Bloque 2']); p = os.path.join(tmp, 'x.xlsx'); wbx.save(p)
        r = run(X, 'import', p)
        check(r.returncode == 1 and needle in r.stdout and open(JP, 'rb').read() == h1, f'import rechaza todo el archivo: {name}', r.stdout)
    wbx = load_workbook(out); wbx['Bloque 2']['K2'] = None; p = os.path.join(tmp, 'x.xlsx'); wbx.save(p)
    run(X, 'import', p)
    check({i['id']: i for i in jl(JP)['incidencias']}[ids2[0]]['estado'] == 'abierta', 'borrar la decisión la deja abierta otra vez')
    # lo aplicado o cerrado no se reabre desde la hoja
    dd = jl(JP)
    for i in dd['incidencias']:
        if i['id'] == ids2[1]: i['estado'] = 'aplicada'
    json.dump(dd, open(JP, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    wbx = load_workbook(out); wbx['Bloque 2']['K3'] = 'Rechazar'; wbx.save(p)
    run(X, 'import', p)
    check({i['id']: i for i in jl(JP)['incidencias']}[ids2[1]]['estado'] == 'aplicada', 'una incidencia aplicada no se reabre desde la hoja')

    # ---- cerrar en bloque y auditoria --v1 ----
    nC = sum(1 for i in jl(JP)['incidencias'] if i['clase'] == 'C' and i['estado'] == 'abierta')
    r = run(X, 'cerrar', 'C', '--dry')
    check(f'{nC} incidencias de clase C cerradas' in r.stdout and sum(1 for i in jl(JP)['incidencias'] if i['estado'] == 'cerrada') == 0, 'cerrar C --dry cuenta y no escribe', r.stdout)
    r = run(X, 'cerrar', 'C')
    cerr = [i for i in jl(JP)['incidencias'] if i['estado'] == 'cerrada']
    check(len(cerr) == nC and all(i['clase'] == 'C' for i in cerr), 'cerrar C cierra solo la clase C abierta', r.stdout)
    check('bloques asignados: 1 (40), 2 (6), 3 (2)' in run(X, 'resumen').stdout, 'resumen por clase y estado con los bloques')
    run('tools/build_data.py')
    P = os.path.join(tmp, 'v1.json')
    run('tools/pipeline/auditoria.py', '--v1', '--json', P)
    abiertas = [i for i in jl(JP)['incidencias'] if i['estado'] in ('abierta', 'decidida')]
    check(len(jl(P)['open_incidents']) == len(abiertas) and len(abiertas) > 0, 'auditoria --v1 cuenta las incidencias abiertas o decididas sin aplicar', len(jl(P)['open_incidents']))
    dd = jl(JP)
    for i in dd['incidencias']:
        if i['estado'] in ('abierta', 'decidida'): i['estado'] = 'cerrada'
    json.dump(dd, open(JP, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    run('tools/pipeline/auditoria.py', '--v1', '--json', P)
    check(jl(P)['open_incidents'] == [], 'con todas cerradas o aplicadas, auditoria --v1 da 0 en incidencias')
    os.remove(JP)
    r = run('tools/pipeline/auditoria.py', '--v1', '--json', P)
    check('falta tools/r2_incidencias.json' in jl(P)['open_incidents'][0], 'sin r2_incidencias.json, auditoria --v1 lo marca como pendiente')
finally:
    shutil.rmtree(tmp, ignore_errors=True)
print('RESULT: ' + ('PASS' if ok_all else 'FAIL'))
sys.exit(0 if ok_all else 1)
